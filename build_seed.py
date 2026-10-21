"""Build supabase/seed.sql (and a demo JSON) from Freezer Map 26.xlsm.

Usage: python3 tools/build_seed.py <workbook.xlsm> <out_dir>
"""
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

SRC, OUT = Path(sys.argv[1]), Path(sys.argv[2])
wb = openpyxl.load_workbook(SRC, data_only=True)

ID_RE = re.compile(r"^[A-Z]{1,2}\d+$")


def d(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    return None


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


# ---------- pallet master list ----------
master = {}
for r in wb["Combined Pallet List"].iter_rows(min_row=2, values_only=True):
    code = str(r[0]).strip().upper() if r[0] is not None else ""
    if not ID_RE.match(code) or code in master:
        continue
    master[code] = dict(
        strain=(str(r[2]).strip() if r[2] not in (None, 0) else None),
        zone=(str(r[3]).strip() if r[3] not in (None, 0) else None),
        og_weight=num(r[4]) or None,
        current_weight=num(r[10]),
        package_tag=(str(r[6]).strip() if r[6] else None),
        harvest_date=d(r[1]),
        planting_date=d(r[7]),
        notes=(r[8].strip() if isinstance(r[8], str) and r[8].strip() else None),
    )

# ---------- maps ----------
freezers, slots, placements = [], [], []  # placements: (slot_id, raw_token)
fid = sid = 0
for sheet, site in (("Freezer Map OHR", "OHR"), ("Freezer Map GV", "GVR")):
    ws = wb[sheet]
    for r in range(1, ws.max_row + 1):
        for col_letter, off in (("B", 0), ("K", 9)):
            v = ws[f"{col_letter}{r}"].value
            if not (isinstance(v, str) and re.match(r"^Freezer \d+$", v.strip())):
                continue
            fid += 1
            status = ws.cell(r, 8 + off).value or "On"
            if status not in ("On", "Off", "Needs Attention", "Not Working"):
                status = "On"
            freezers.append(dict(id=fid, site=site, number=int(v.split()[1]), status=status))
            for i in range(9):
                rr = r + 2 + i
                for colno, (pc, sc) in enumerate(((2 + off, 5 + off), (6 + off, 9 + off)), 1):
                    sid += 1
                    code = ws.cell(rr, sc).value or f"{freezers[-1]['number']}{'ABCDEFGHIJKLMNOPQR'[(colno-1)*9+i]}"
                    slots.append(dict(id=sid, freezer_id=fid, code=str(code), col=colno,
                                      row=i + 1, front_rank=9 - i))
                    cell = ws.cell(rr, pc).value
                    if cell not in (None, ""):
                        placements.append((sid, str(cell).strip()))

slot_by_id = {s["id"]: s for s in slots}
fz_by_id = {f["id"]: f for f in freezers}


def where(slot_id):
    s = slot_by_id[slot_id]
    return f"{fz_by_id[s['freezer_id']]['site']} {s['code']}"


# ---------- pallets ----------
pallets = {}  # code -> row
issues = []
for slot_id, raw in placements:
    tokens = [t.strip() for t in raw.split("/") if t.strip()]
    for tok in tokens:
        code = tok.upper()
        if not ID_RE.match(code):
            code = f"X-{where(slot_id).replace(' ', '-')}"
            n = 2
            while code in pallets:
                code = f"X-{where(slot_id).replace(' ', '-')}-{n}"; n += 1
            pallets[code] = dict(code=code, label=tok, slot_id=slot_id, status="in_freezer",
                                 needs_review=f'Not a pallet number: "{tok}"')
            continue
        if code in pallets:
            p = pallets[code]
            msg = f"Also listed at {where(slot_id)}"
            p["needs_review"] = (p.get("needs_review") + "; " if p.get("needs_review") else "") + msg
            continue
        m = master.get(code)
        row = dict(code=code, slot_id=slot_id, status="in_freezer")
        if tok != code:
            row["label"] = tok
        if m is None:
            row["needs_review"] = "Not in pallet list"
        else:
            row.update(m)
            if not m["strain"] or not m["current_weight"]:
                row["needs_review"] = "Missing strain or weight in pallet list"
        pallets[code] = row

# pallets in the list but not on a map: keep for history/search
for code, m in master.items():
    if code in pallets:
        continue
    if not m["strain"] and not m["current_weight"]:
        continue  # unused placeholder rows
    pallets[code] = dict(code=code, status="archived", slot_id=None, **m)

plist = list(pallets.values())
for i, p in enumerate(plist, 1):
    p["id"] = i

# ---------- SQL ----------
def q(v):
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return repr(round(v, 2))
    return "'" + str(v).replace("'", "''") + "'"


cols = ["id", "code", "label", "strain", "zone", "harvest_date", "planting_date", "og_weight",
        "current_weight", "package_tag", "status", "slot_id", "needs_review", "notes"]
lines = ["-- Seed data generated from Freezer Map 26.xlsm", "begin;",
         "truncate events, pallets, slots, freezers restart identity cascade;"]
lines.append("insert into freezers (id, site, number, status) values")
lines.append(",\n".join(f"({f['id']},{q(f['site'])},{f['number']},{q(f['status'])})" for f in freezers) + ";")
lines.append("insert into slots (id, freezer_id, code, col, row, front_rank) values")
lines.append(",\n".join(f"({s['id']},{s['freezer_id']},{q(s['code'])},{s['col']},{s['row']},{s['front_rank']})" for s in slots) + ";")
for chunk in range(0, len(plist), 500):
    part = plist[chunk:chunk + 500]
    lines.append(f"insert into pallets ({', '.join(cols)}) values")
    lines.append(",\n".join("(" + ",".join(q(p.get(c)) for c in cols) + ")" for p in part) + ";")
lines.append("insert into events (user_email, type, pallet_id, to_slot_id, lbs, note) "
             "select 'import', 'import', id, slot_id, current_weight, 'Imported from Freezer Map 26.xlsm' "
             "from pallets where status = 'in_freezer';")
for t in ("freezers", "slots", "pallets"):
    lines.append(f"select setval(pg_get_serial_sequence('{t}','id'), (select max(id) from {t}));")
lines.append("commit;")
(OUT / "supabase").mkdir(parents=True, exist_ok=True)
(OUT / "supabase" / "seed.sql").write_text("\n".join(lines) + "\n")

# demo data for the offline preview
(OUT / "web" / "demo").mkdir(parents=True, exist_ok=True)
(OUT / "web" / "demo" / "demo-data.json").write_text(json.dumps(dict(
    freezers=freezers, slots=slots,
    pallets=[{c: p.get(c) for c in cols} for p in plist if p["status"] != "archived"])))

placed = [p for p in plist if p["status"] == "in_freezer"]
review = [p for p in plist if p.get("needs_review")]
print(f"freezers {len(freezers)}  slots {len(slots)}  pallets {len(plist)} "
      f"(in freezers {len(placed)}, archived {len(plist)-len(placed)})  needs review {len(review)}")
for p in review[:400]:
    print("  ", p["code"], "|", p.get("label") or "", "|", where(p["slot_id"]) if p.get("slot_id") else "", "|", p["needs_review"])
