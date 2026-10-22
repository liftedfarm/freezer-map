"""Read the three cultivation workbooks into one JSON file (input for build_cult_import.py).

Usage: python3 -I tools/cult_extract.py <Harvest_Scheduling_Workbook.xlsx> <Farm Data OHR 2026.xlsx> <2026 Strain Selections.xlsx> <out.json>
"""
import datetime as dt
import json
import re
import sys
import warnings

import openpyxl
from openpyxl.utils import column_index_from_string as CI

warnings.filterwarnings('ignore')
hsw_f, farm_f, sel_f, out_f = sys.argv[1:5]
books = {k: openpyxl.load_workbook(f, data_only=True, read_only=True) for k, f in (('hsw', hsw_f), ('farm', farm_f), ('sel', sel_f))}


def clean(v):
    if isinstance(v, str):
        v = re.sub(r'\s+', ' ', v.replace('\xa0', ' ')).strip()
        return v or None
    return v


def num(v):
    v = clean(v)
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    if not isinstance(v, str):
        return None
    m = re.fullmatch(r'[-+]?[\d,]*\.?\d+', v)
    return float(v.replace(',', '')) if m else None


def date(v):
    v = clean(v)
    if isinstance(v, dt.datetime):
        return None if v.year < 2000 else v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, str):
        m = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})', v)
        if m:
            return dt.date(int(m[3]), int(m[1]), int(m[2])).isoformat()
    return None


def rows(book, sheet, cols, start, need):
    ws = books[book][sheet]
    out = []
    for i, r in enumerate(ws.iter_rows(min_row=start, values_only=True), start=start):
        def get(c):
            j = CI(c) - 1
            return r[j] if j < len(r) else None
        if clean(get(need)) in (None, 0):
            continue
        d = {k: get(c) for k, c in cols.items()}
        d['_row'] = i
        out.append(d)
    return out


out = {}

# ---- strain lists
out['variety_lists'] = []
for col, kind in (('A', 'OHR'), ('B', 'GVR'), ('D', 'OHR biomass')):
    for r in rows('hsw', 'Variety Lists', {'name': col}, 2, col):
        out['variety_lists'].append({'name': clean(r['name']), 'list': kind})
out['labs'] = [clean(r['name']) for r in rows('hsw', 'Variety Lists', {'name': 'F'}, 2, 'F')]

# ---- harvest log (OHR, GVR)
HCOLS = dict(strain='A', field='B', type='C', date='D', plants='E', lpp='F', rawpp='G', deleaf='H', buck='I', conv='J', notes='V')
out['harvest_log'] = []
for site in ('OHR', 'GVR'):
    for r in rows('hsw', site, HCOLS, 9, 'A'):
        t = clean(r['type']) or ''
        out['harvest_log'].append(dict(site=site, strain=clean(r['strain']), field=str(clean(r['field']) or ''),
            kind='dry' if t.lower().startswith('dry') else 'ff', date=date(r['date']), plants=num(r['plants']),
            lpp=num(r['lpp']), rawpp=num(r['rawpp']), deleaf=num(r['deleaf']), buck=num(r['buck']), conv=num(r['conv']),
            notes=clean(r['notes']), row=r['_row']))
out['site_params'] = {'OHR': {}, 'GVR': {}}
for site in ('OHR', 'GVR'):
    for r in rows('hsw', site, {'label': 'A', 'v': 'G'}, 4, 'A')[:3]:
        out['site_params'][site][clean(r['label'])] = num(r['v'])

out['biomass_log'] = [dict(strain=clean(r['strain']), field=clean(r['field']), acres=num(r['acres']), ypa=num(r['ypa']),
                          method=clean(r['method']), cut=date(r['cut']), merge=num(r['merge']), harvest=num(r['harvest']),
                          notes=clean(r['notes']), row=r['_row'])
                      for r in rows('hsw', 'OHR Biomass', dict(strain='A', field='B', acres='C', ypa='D', method='E', cut='F', merge='G', harvest='H', notes='L'), 9, 'A')]

# ---- testing entries
out['area_tests'] = [dict(site=clean(r['site']), field=str(clean(r['field'])), lab=clean(r['lab']), submitted=date(r['sub']),
                          results=date(r['res']), result=clean(r['pf']), notes=clean(r['notes']))
                     for r in rows('hsw', 'Pre-Harvest Testing Schedule', dict(site='A', field='B', lab='J', sub='K', res='L', pf='M', notes='N'), 12, 'A')
                     if r['_row'] <= 23 and clean(r['site']) in ('OHR', 'GVR')]
# strain tests: the entries in J:N belong to the Test ID in column Z (the sheet's rows shift when the schedule changes)
st = rows('hsw', 'Pre-Harvest Testing Schedule', dict(strain='A', site='B', field='C', date='D', tid='I', lab='J', sub='K', res='L', pf='M', notes='N', snap='Z'), 28, 'A')
out['strain_tests'] = []
for r in st:
    key = clean(r['snap']) or clean(r['tid'])
    if not key or not any(clean(r[k]) for k in ('lab', 'sub', 'res', 'pf', 'notes')):
        continue
    m = re.match(r'(.+) \| (OHR|GVR) (.+)', key)
    if m:
        out['strain_tests'].append(dict(strain=m[1], site=m[2], field=m[3], lab=clean(r['lab']), submitted=date(r['sub']),
                                        results=date(r['res']), result=clean(r['pf']), notes=clean(r['notes'])))

# ---- flower selections
out['flower_picks'] = [dict(strain=clean(r['s']), lbs=num(r['lbs'])) for r in rows('hsw', 'Flower Selections', dict(s='B', lbs='C'), 9, 'B')
                       if num(r['lbs']) and clean(r['s']) != 'Total']

# ---- plantings (Farm Data OHR)
out['ohr_plantings'] = [dict(strain=clean(r['strain']), batch=num(r['batch']), zone=clean(r['zone']), license=clean(r['lic']),
                             loss=num(r['loss']), planted=date(r['planted']), beds=num(r['beds']), spacing=num(r['sp']),
                             plants=num(r['plants']), adj=clean(r['adj']), testing=clean(r['testing']), first_flower=date(r['ff']),
                             notes=clean(r['notes']), row=r['_row'])
                        for r in rows('farm', 'Clone Planting', dict(strain='B', batch='C', zone='D', lic='E', loss='F', planted='H', beds='I', sp='J', plants='K', adj='M', testing='N', ff='O', notes='S'), 3, 'B')]
out['ohr_biomass'] = []
for r in rows('farm', 'Biomass #1', dict(strain='B', batch='C', zone='D', acres='E', lic='F', seed='G', plants='P', total='J', planted='L', sp='N', notes='S'), 3, 'B'):
    out['ohr_biomass'].append(dict(strain=clean(r['strain']), batch=num(r['batch']), zone=clean(r['zone']), acres=num(r['acres']), license=clean(r['lic']),
                                   seed_lbs=num(r['seed']), plants=num(r['plants']) or num(r['total']), planted=date(r['planted']), spacing=num(r['sp']),
                                   first_flower=None, notes=clean(r['notes']), src='Biomass #1'))
for r in rows('farm', 'Biomass #2', dict(strain='B', batch='C', zone='D', lic='E', seed='F', plants='M', planted='J', sp='L', ff='P', notes='Q'), 3, 'B'):
    out['ohr_biomass'].append(dict(strain=clean(r['strain']), batch=num(r['batch']), zone=clean(r['zone']), acres=None, license=clean(r['lic']),
                                   seed_lbs=num(r['seed']), plants=num(r['plants']), planted=date(r['planted']), spacing=num(r['sp']),
                                   first_flower=date(r['ff']), notes=clean(r['notes']), src='Biomass #2'))
out['receiving'] = [dict(strain=clean(r['s']), batch=num(r['b']), received=num(r['recv']), rejected=num(r['rej']), nursery=clean(r['n']),
                         source=clean(r['src']), date=date(r['d']))
                    for r in rows('farm', 'Plant Receiving', dict(s='B', b='C', recv='D', rej='E', n='K', src='L', d='M'), 2, 'B') if num(r['recv'])]

# ---- harvest actuals (Farm Data OHR)
out['metrc'] = [dict(strain=clean(r['s']), date=date(r['d']), zone=clean(r['z']), type=clean(r['t']), plants=num(r['p']) or num(r['p2']),
                     raw=num(r['raw']), frozen=num(r['fz']), packages=num(r['n']), name=clean(r['name']))
                for r in rows('farm', 'Harvest METRC', dict(s='A', d='B', z='C', t='D', p='E', raw='F', p2='G', fz='H', n='K', name='L'), 3, 'A') if date(r['d'])]
out['dry_harvests'] = [dict(strain=clean(r['s']), date=date(r['d']), zone=clean(r['z']), plants=num(r['p']), raw=num(r['raw']),
                            tag=clean(r['tag']), bucked=num(r['b']), trim=num(r['trim']), mold=num(r['mold']), notes=clean(r['notes']))
                       for r in rows('farm', 'Clone Dry Harvest ', dict(s='A', d='B', z='C', p='E', raw='F', tag='G', b='H', trim='I', mold='J', notes='K'), 3, 'A') if date(r['d'])]
ws = books['farm']['Raw Weights']
grid = [list(r) for r in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True)]
out['raw_weights'] = []
for c in range(1, len(grid[0]), 2):
    s = clean(grid[0][c]) if c < len(grid[0]) else None
    if not s:
        continue
    bins = [num(grid[r][c]) for r in range(7, len(grid)) if c < len(grid[r]) and num(grid[r][c])]
    out['raw_weights'].append(dict(strain=s, date=date(grid[1][c]), zone=clean(grid[2][c]), type=clean(grid[3][c]),
                                   plants=num(grid[4][c]) or num(grid[6][c]), total=num(grid[5][c]), bins=bins))

# ---- Data Tapes (Strain Selections): plantings with projections + actuals
TAPE = dict(strain='C', roshi='D', item='E', farm='F', field='G', nursery='I', type='J', qty='K', planted='AM', sp='AN', acres='AO',
            acres_act='AP', pdate='AQ', lpp='AT', ffd='AX', ffp='AY', fflbs='AZ', dryd='BC', dryp='BD', equiv='BE')
out['tape_2026'] = []
for sheet, site in (('OHR 2026 Data Tape', 'OHR'), ('GVR 2026 Data Tape', 'GVR')):
    for r in rows('sel', sheet, TAPE, 4, 'C'):
        if not num(r['planted']) or not clean(r['field']) or str(clean(r['strain'])).startswith('Total'):
            continue
        out['tape_2026'].append(dict(site=site, strain=clean(r['strain']), roshi=clean(r['roshi']), item=clean(r['item']), field=str(clean(r['field']) or ''),
            nursery=clean(r['nursery']), type=clean(r['type']), plants=num(r['planted']), spacing=num(r['sp']), acres=num(r['acres_act']) or num(r['acres']),
            planted=date(r['pdate']), lpp=num(r['lpp']), ff_date=date(r['ffd']), ff_plants=num(r['ffp']), ff_lbs=num(r['fflbs']),
            dry_date=date(r['dryd']), dry_plants=num(r['dryp'])))
out['gvr_plantings'] = [dict(strain=clean(r['s']), zone=str(clean(r['z']) or ''), license=clean(r['lic']), loss=(num(r['l1']) or 0) + (num(r['l2']) or 0),
                             replaced=num(r['rep']), planted=date(r['d']), beds=num(r['beds']), spacing_ft=num(r['sp']), plants=num(r['p']),
                             acres=num(r['a']), first_flower=date(r['ff']), notes=' '.join(x for x in (clean(r['n1']), clean(r['n2'])) if x) or None)
                        for r in rows('sel', 'GVPlanting', dict(s='B', z='C', lic='D', l1='E', l2='F', rep='H', d='I', beds='J', sp='K', p='L', a='M', ff='N', n1='Q', n2='R'), 2, 'B')
                        if num(r['p'])]
# older seasons
out['history'] = []
for r in rows('sel', '2025 Data Tape', dict(strain='C', item='D', farm='E', nursery='H', planted='AC', pdate='AG', ffd='AQ', ffp='AR', fflbs='AS', equiv='AZ'), 7, 'C'):
    if num(r['ffp']) and num(r['fflbs']):
        out['history'].append(dict(year=2025, farm=clean(r['farm']), strain=clean(r['strain']), plants=num(r['ffp']), lbs=num(r['fflbs']),
                                   planted=date(r['pdate']), harvested=date(r['ffd'])))
ws = books['sel']['2024 Data Tape']
hdr = None
for i, r in enumerate(ws.iter_rows(min_row=1, max_row=12, values_only=True), start=1):
    labels = [clean(x) for x in r]
    if 'Harvested FF LBs' in labels or any(isinstance(x, str) and 'Harvested' in x and 'FF' in x for x in labels):
        hdr = (i, labels)
        break
out['history_2024_header'] = hdr[1] if hdr else None
if hdr:
    lab = hdr[1]
    def idx(pred):
        for j, x in enumerate(lab):
            if isinstance(x, str) and pred(x):
                return j
    j_strain = idx(lambda x: x == 'Strain')
    j_farm = idx(lambda x: x == 'Farm')
    j_hd = idx(lambda x: x == 'Harvest Date')
    j_p = idx(lambda x: x.startswith('# FF Plants'))
    j_lbs = idx(lambda x: x.startswith('Harvested FF'))
    for r in ws.iter_rows(min_row=hdr[0] + 1, values_only=True):
        if j_strain is None or j_lbs is None:
            break
        s, p, l = clean(r[j_strain]), num(r[j_p]) if j_p is not None else None, num(r[j_lbs])
        if s and p and l:
            out['history'].append(dict(year=2024, farm=clean(r[j_farm]) if j_farm is not None else None, strain=s, plants=p, lbs=l,
                                       planted=None, harvested=date(r[j_hd]) if j_hd is not None else None))

# ---- 2026 crop plan, trials, nursery orders
out['crop_plan'] = [dict(strain=clean(r['clean']) or clean(r['s']), nursery=clean(r['n']), farm=clean(r['farm']), plants=num(r['q']), lpp=num(r['y']),
                         spacing=num(r['sp']), d1=num(r['d1']), d1date=date(r['d1d']), p1date=date(r['p1d']), d2=num(r['d2']), d2date=date(r['d2d']),
                         p2date=date(r['p2d']), notes=clean(r['notes']))
                    for r in rows('sel', 'Clone Crop Plan', dict(s='B', clean='C', n='D', farm='E', q='F', y='H', sp='I', d1='L', d1d='M', p1d='N', d2='O', d2d='P', p2d='Q', notes='R'), 2, 'B')
                    if num(r['q'])]
out['trials'] = [dict(strain=clean(r['s']), nursery=clean(r['n']), farm=clean(r['farm']), plants=num(r['q']), lpp=num(r['y']), spacing=num(r['sp']))
                 for r in rows('sel', 'Field Trials', dict(s='A', n='B', farm='C', q='D', y='F', sp='G'), 2, 'A') if num(r['q'])]

json.dump(out, open(out_f, 'w'), indent=1, default=str)
print({k: len(v) if isinstance(v, (list, dict)) else v for k, v in out.items()})
